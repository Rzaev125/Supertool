"""Capture the actual profile policy used for DOCX, XLSX and PPTX engine launches."""
import argparse
import hashlib
from pathlib import Path
import shutil
import subprocess
import tempfile
from docx import Document
from pptx import Presentation
import xlsxwriter

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--exe', required=True, type=Path)
parser.add_argument('--probe', required=True, type=Path)
parser.add_argument('--cancel-probe', type=Path)
args = parser.parse_args()
parent = Path.cwd() / 'test-output'
parent.mkdir(exist_ok=True)
root = Path(tempfile.mkdtemp(prefix='office-policy-', dir=parent))
try:
    app = root / 'ShiftWheel.exe'
    shutil.copy2(args.exe, app)
    engine = root / 'modules/libreoffice/program/soffice.com'
    engine.parent.mkdir(parents=True)
    shutil.copy2(args.probe, engine)
    doc = root / 'document.docx'
    document = Document(); document.add_paragraph('Profile policy fixture'); document.save(doc)
    book = root / 'workbook.xlsx'
    with xlsxwriter.Workbook(book) as workbook:
        workbook.add_worksheet().write(0, 0, 42)
    deck = root / 'slides.pptx'
    presentation = Presentation(); presentation.slides.add_slide(presentation.slide_layouts[6]); presentation.save(deck)
    for source in [doc, book, deck]:
        before = hashlib.sha256(source.read_bytes()).digest()
        destination, work = root / (source.stem + '.pdf'), root / ('work-' + source.stem)
        process = subprocess.run([str(app), '--convert', str(source), str(destination), str(work), 'pdf'],
                                 capture_output=True, timeout=30, creationflags=subprocess.CREATE_NO_WINDOW)
        error = Path(str(work) + '.error.txt')
        assert process.returncode == 0, error.read_text('utf-8') if error.exists() else process.stderr
        assert destination.read_bytes() == b'%PDF-policy-probe-only\n'
        assert hashlib.sha256(source.read_bytes()).digest() == before and not work.exists()
    print('PASS actual DOCX/XLSX/PPTX engine profiles: OpenCL disabled, macros isolated, originals preserved', flush=True)
    if args.cancel_probe:
        result = subprocess.run([str(args.cancel_probe), str(doc)], capture_output=True, timeout=25,
                                creationflags=subprocess.CREATE_NO_WINDOW)
        assert result.returncode == 0, result.stdout + result.stderr
        print(result.stdout.decode('utf-8').strip(), flush=True)
except BaseException:
    print('Office policy fixture retained: ' + str(root), flush=True)
    raise
else:
    assert root.resolve().parent == parent.resolve()
    shutil.rmtree(root)
