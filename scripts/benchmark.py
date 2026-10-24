"""Measure local builds on identical fixtures; never fetch old releases."""
import argparse
import hashlib
import json
import pathlib
import shutil
import statistics
import subprocess
import tempfile
from PIL import Image, ImageChops, ImageDraw, ImageStat
from docx import Document
from pptx import Presentation
from reportlab.pdfgen import canvas
from pypdf import PdfReader
import xlsxwriter

ROOT = pathlib.Path(__file__).resolve().parents[1]


def sha(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def run(arguments):
    return subprocess.run([str(a) for a in arguments], check=True, capture_output=True,
                          timeout=330, creationflags=subprocess.CREATE_NO_WINDOW)


def summarize(rows):
    result = {}
    for case in dict.fromkeys(r['case'] for r in rows):
        result[case] = {}
        for version in dict.fromkeys(r['version'] for r in rows if r['case'] == case):
            values = sorted((r for r in rows if r['case'] == case and r['version'] == version), key=lambda r:r['iteration'])
            repeats = values[1:]
            result[case][version] = dict(runs=len(values), first_wall_ms=values[0]['wall_ms'],
                repeat_median_wall_ms=statistics.median(r['wall_ms'] for r in repeats) if repeats else None,
                median_wall_ms=statistics.median(r['wall_ms'] for r in values),
                median_cpu_ms=statistics.median(r['cpu_ms'] for r in values),
                cpu_one_core_percent=100 * sum(r['cpu_ms'] for r in values) / sum(r['wall_ms'] for r in values),
                peak_commit_bytes=max(r['peak_commit_bytes'] for r in values),
                sampled_sum_rss_bytes=max(r['sampled_sum_rss_bytes'] for r in values))
    return result


def compare_all_outputs(outputs, extension, executable, work):
    reference = next(iter(outputs.values()))[0]
    if extension != 'pdf':
        expected = reference.read_bytes()
        for paths in outputs.values():
            for path in paths:
                assert path.read_bytes() == expected, 'Result changed: ' + str(path)
        return
    expected = PdfReader(reference)
    expected_text = [page.extract_text() for page in expected.pages]
    expected_boxes = [tuple(page.mediabox) for page in expected.pages]
    rendered = []
    for paths in outputs.values():
        for path in paths:
            reader = PdfReader(path)
            assert [page.extract_text() for page in reader.pages] == expected_text, 'PDF text changed'
            assert [tuple(page.mediabox) for page in reader.pages] == expected_boxes, 'PDF page geometry changed'
            # Rendering is outside the timed interval. Compare decoded pixels,
            # ignoring nondeterministic PDF creation timestamps and internal IDs.
            folder, staging = work / (path.stem + '-pages'), work / (path.stem + '-render-work')
            run([executable, '--convert', path, folder, staging, 'png'])
            pixels = []
            for image in sorted(folder.glob('*.png')):
                with Image.open(image) as decoded:
                    decoded.load(); pixels.append((decoded.size, decoded.convert('RGBA').tobytes()))
            assert len(pixels) == len(expected.pages), 'PDF rendering omitted pages'
            rendered.append(pixels)
    assert all(value == rendered[0] for value in rendered), 'PDF page pixels changed'


def fixtures(work, ffmpeg):
    options = work / 'options.txt'
    options.write_text('version=1\nquality=balanced\nresources=balanced\nvideo_acceleration=cpu\npdf_ocr=1\nocr_languages=eng\n',encoding='ascii')
    image = Image.new('RGB',(1600,1000),'#c9e3ef'); draw = ImageDraw.Draw(image)
    for i in range(0,1600,20): draw.rectangle((i,i%700,i+10,900),fill=(i%255,90,160))
    image.save(work/'image.png')
    run([ffmpeg,'-hide_banner','-loglevel','error','-y','-f','lavfi','-i','testsrc2=size=1280x720:rate=30','-t','4','-c:v','libx264','-threads','2','-pix_fmt','yuv420p',work/'video.mp4'])
    run([ffmpeg,'-hide_banner','-loglevel','error','-y','-f','lavfi','-i','sine=frequency=440:duration=10',work/'audio.wav'])
    document = Document(); document.add_heading('Supertool benchmark',0)
    document.add_paragraph('Identical source. Unicode: Привет, Bakı.'); document.save(work/'document.docx')
    with xlsxwriter.Workbook(work/'workbook.xlsx') as workbook:
        for name,value in [('First',42),('Second',73)]:
            sheet=workbook.add_worksheet(name); sheet.write_row(0,0,[name,value]); sheet.write_formula(1,1,'=B1+1',None,value+1)
    deck=Presentation()
    for title in ['Slide one','Slide two']:
        slide=deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text=title
    deck.save(work/'slides.pptx')
    pdf=canvas.Canvas(str(work/'text.pdf')); pdf.drawString(72,720,'Supertool deterministic text fixture'); pdf.save()
    scan=Image.new('RGB',(1500,700),'white')
    from PIL import ImageFont
    ImageDraw.Draw(scan).text((90,200),'SHIFTWHEEL OCR BENCHMARK 12345',fill='black',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',60))
    scan.save(work/'scan.pdf','PDF',resolution=144)
    return options


def inspect_output(source, output, extension, ffmpeg):
    if extension == 'pdf':
        reader=PdfReader(output); text=' '.join(p.extract_text() for p in reader.pages)
        if source.suffix == '.docx': assert len(reader.pages)==1 and 'Identical source' in text
        if source.suffix == '.xlsx': assert len(reader.pages)==2 and all(s in text for s in ['42','43','73','74'])
        if source.suffix == '.pptx': assert len(reader.pages)==2 and 'Slide one' in text and 'Slide two' in text
    elif extension == 'jpg':
        with Image.open(source) as original, Image.open(output) as result:
            result.load(); assert result.format=='JPEG' and result.size==original.size
            assert max(ImageStat.Stat(ImageChops.difference(original.convert('RGB'),result.convert('RGB'))).mean)<12
    elif extension in ('mp4','flac'):
        probe=json.loads(run([ffmpeg.with_name('ffprobe.exe'),'-v','error','-show_streams','-show_format','-of','json',output]).stdout)
        assert abs(float(probe['format']['duration'])-(4 if extension=='mp4' else 10))<0.15
        if extension=='mp4': assert any(s['codec_type']=='video' and s['width']==1280 and s['height']==720 for s in probe['streams'])
        run([ffmpeg,'-v','error','-i',output,'-f','null','-'])
        if extension=='flac':
            def pcm(path): return run([ffmpeg,'-v','error','-i',path,'-f','s16le','-acodec','pcm_s16le','-']).stdout
            assert pcm(source)==pcm(output), 'Lossless audio samples changed'
    else:
        text=' '.join(output.read_text('utf-8').split())
        assert ('SHIFTWHEEL OCR BENCHMARK 12345' in text.upper()) if source.name=='scan.pdf' else ('Supertool deterministic text fixture' in text)


def execute(work, executable, baseline, iterations):
    runner=ROOT/'build/Release/benchmark_runner.exe'; ffmpeg=ROOT/'tools/modules/ffmpeg/ffmpeg.exe'
    options=fixtures(work,ffmpeg)
    cases=[('image.png','jpg'),('video.mp4','mp4'),('audio.wav','flac'),('document.docx','pdf'),
           ('workbook.xlsx','pdf'),('slides.pptx','pdf'),('text.pdf','txt'),('scan.pdf','txt')]
    originals={name:sha(work/name) for name,_ in cases}; rows=[]
    programs=([('baseline',baseline)] if baseline else [])+[('current',executable)]
    for index,(name,extension) in enumerate(cases):
        outputs={version:[] for version,_ in programs}
        for iteration in range(iterations):
            order=programs if iteration%2==0 else list(reversed(programs))
            for version,program in order:
                tag=f'{index}-{version}-{iteration}'; output=work/f'{tag}.{extension}'; staging=work/f'{tag}.tmp'
                metrics=json.loads(run([runner,program,'--convert',work/name,output,staging,'--options',options]).stdout)
                assert metrics['exit']==0 and output.is_file() and not staging.exists()
                inspect_output(work/name,output,extension,ffmpeg)
                metrics.update(case=name,version=version,iteration=iteration,phase='first_in_series' if iteration==0 else 'repeat',
                               output_bytes=output.stat().st_size,output_sha256=sha(output))
                rows.append(metrics);outputs[version].append(output)
        compare_all_outputs(outputs,extension,executable,work)
        assert originals[name]==sha(work/name), 'Original changed: '+name
        print(name,'verified (every result, identical content)',flush=True)
    report=dict(environment='Windows x64; sequential; balanced; CPU video; synthetic fixtures; no general performance guarantee',
        interpretation='First in series is not OS-cold. Every conversion uses a fresh process and Office profile. No OS caches purged. CPU % is one-core equivalent; RSS sums shared pages across processes and is sampled every ~5ms; commit is PeakJobMemoryUsed. Validation is outside timed intervals.',
        iterations=iterations,executables={k:str(v) for k,v in programs},executable_sha256={k:sha(v) for k,v in programs},
        inputs=originals,options_snapshot=options.read_text('ascii'),options_sha256=sha(options),
        script_sha256=sha(pathlib.Path(__file__)),rows=rows,summary=summarize(rows),content_verified=True)
    (ROOT/'build/benchmark-results.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    for name,versions in report['summary'].items():
        print(name,[(v,round(s['median_wall_ms'],1),round(s['peak_commit_bytes']/2**20,1)) for v,s in versions.items()],flush=True)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--exe',type=pathlib.Path,default=ROOT/'build/Release/Supertool.exe')
    parser.add_argument('--baseline',type=pathlib.Path,help='Optional local EXE with matching modules beside it.')
    parser.add_argument('--iterations',type=int,choices=range(1,11),default=3)
    args=parser.parse_args()
    work=pathlib.Path(tempfile.mkdtemp(prefix='benchmark-',dir=ROOT/'build'))
    try:
        execute(work,args.exe.resolve(),args.baseline.resolve() if args.baseline else None,args.iterations)
    except BaseException:
        print('Failed benchmark fixtures retained: '+str(work),flush=True)
        raise
    else:
        assert work.resolve().parent == (ROOT/'build').resolve()
        shutil.rmtree(work)

if __name__=='__main__': main()
