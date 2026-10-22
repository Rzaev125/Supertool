"""Real presentation imports, all-slide exports, and refusal/preservation checks."""
import argparse
import hashlib
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pypdf import PdfReader
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')
sys.stderr.reconfigure(encoding='utf-8')

parser = argparse.ArgumentParser()
parser.add_argument('--exe', required=True, type=Path)
args = parser.parse_args()
exe = args.exe.resolve()
parent = Path.cwd() / 'test-output'
parent.mkdir(exist_ok=True)
root = Path(tempfile.mkdtemp(prefix='presentations-', dir=parent)) / 'Слайды с пробелами'
root.mkdir()
counter = 0


def run(command, timeout=150):
    return subprocess.run([str(x) for x in command], capture_output=True, timeout=timeout,
                          creationflags=subprocess.CREATE_NO_WINDOW)


def digest(path):
    return hashlib.sha256(path.read_bytes()).digest()


def convert(source, target, fail=False, destination=None):
    global counter
    counter += 1
    destination = destination or root / (f'output-{counter}' if target in ['png', 'jpg'] else f'output-{counter}.{target}')
    work = root / f'work-{counter}'
    before = digest(source)
    result = run([exe, '--convert', source, destination, work, target])
    error = Path(str(work) + '.error.txt')
    detail = error.read_text('utf-8') if error.exists() else result.stderr.decode('utf-8', errors='replace')
    assert digest(source) == before, 'presentation source changed'
    assert not work.exists(), 'presentation staging leaked'
    if fail:
        assert result.returncode != 0 and not destination.exists(), 'invalid presentation published output'
    else:
        assert result.returncode == 0 and destination.exists(), detail
    return destination


source = root / 'Презентация.pptx'
deck = Presentation()
deck.slide_width, deck.slide_height = Inches(8), Inches(4.5)
for i, color in enumerate([(220, 30, 40), (20, 50, 220)]):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*color)
    box = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(6), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = ['First slide', 'Second hidden slide'][i]
    paragraph.font.size = Pt(32)
    paragraph.font.color.rgb = RGBColor(255, 255, 255)
    paragraph.font.name = 'Arial'
    if i == 1:
        slide._element.set('show', '0')
deck.save(source)

pdf = convert(source, 'pdf')
pages = PdfReader(pdf).pages
assert len(pages) == 2, 'hidden or later slide omitted'
assert 'First slide' in pages[0].extract_text() and 'Second hidden slide' in pages[1].extract_text(), 'slide order/text lost'
assert abs(float(pages[0].mediabox.width) - 576) < 1 and abs(float(pages[0].mediabox.height) - 324) < 1, 'slide aspect lost'
for target in ['png', 'jpg']:
    folder = convert(source, target)
    images = sorted(folder.glob(f'*.{target}'))
    assert len(images) == 2, 'not all slides rendered'
    for path, expected in zip(images, [(220, 30, 40), (20, 50, 220)]):
        with Image.open(path) as img:
            assert img.size == (1152, 648), 'slide DPI/dimensions ignored'
            actual = img.convert('RGB').getpixel((30, 30))
            assert all(abs(a - b) <= 4 for a, b in zip(actual, expected)), 'slide order/background changed'
print('PASS presentations: all PDF/PNG/JPG slides including hidden, order, text, aspect and pixels', flush=True)

rewritten = convert(source, 'pptx')
assert len(Presentation(rewritten).slides) == 2, 'PPTX export lost a slide'
odp = convert(source, 'odp')
with zipfile.ZipFile(odp) as archive:
    assert archive.read('mimetype') == b'application/vnd.oasis.opendocument.presentation'
roundtrip = convert(odp, 'pptx')
assert len(Presentation(roundtrip).slides) == 2, 'ODP roundtrip lost a slide'
roundtrip_pdf = convert(odp, 'pdf')
assert len(PdfReader(roundtrip_pdf).pages) == 2, 'ODP import lost a slide'

# Produce a legacy PPT fixture with the same installed converter, then exercise the app import.
legacy_dir = root / 'legacy'; legacy_dir.mkdir()
profile = root / 'legacy-profile'
lo = exe.parent / 'modules/libreoffice/program/soffice.com'
result = run([lo, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--nologo', '--nodefault', '--norestore',
              '--convert-to', 'ppt:MS PowerPoint 97', '--outdir', legacy_dir, source])
legacy = legacy_dir / (source.stem + '.ppt')
assert result.returncode == 0 and legacy.exists(), 'legacy PPT fixture generation failed'
legacy_pdf = convert(legacy, 'pdf')
assert len(PdfReader(legacy_pdf).pages) == 2, 'legacy PPT import lost a slide'
print('PASS presentations: PPTX export, ODP roundtrip, legacy PPT input', flush=True)

corrupt = root / 'broken.pptx'; corrupt.write_bytes(b'not an archive')
convert(corrupt, 'pdf', fail=True)
disguised = root / 'word.pptx'
with zipfile.ZipFile(disguised, 'w') as archive:
    archive.writestr('[Content_Types].xml', '<Types/>')
    archive.writestr('word/document.xml', '<document/>')
convert(disguised, 'pdf', fail=True)
macro = root / 'macro.pptx'
with zipfile.ZipFile(source) as old, zipfile.ZipFile(macro, 'w') as new:
    for item in old.infolist():
        new.writestr(item, old.read(item.filename))
    new.writestr('ppt/vbaProject.bin', b'macro fixture must never load')
convert(macro, 'pdf', fail=True)
macro_odp = root / 'macro.odp'
with zipfile.ZipFile(odp) as old, zipfile.ZipFile(macro_odp, 'w') as new:
    for item in old.infolist():
        new.writestr(item, old.read(item.filename))
    new.writestr('Basic/Standard/Module1.xml', '<module/>')
convert(macro_odp, 'pdf', fail=True)

existing = root / 'existing.pdf'; existing.write_bytes(b'foreign result')
work = root / 'collision-work'; before = digest(source)
result = run([exe, '--convert', source, existing, work, 'pdf'])
assert result.returncode != 0 and existing.read_bytes() == b'foreign result' and digest(source) == before and not work.exists()
foreign = root / 'foreign-work'; foreign.mkdir(); (foreign / 'keep').write_bytes(b'keep')
destination = root / 'foreign.pdf'
result = run([exe, '--convert', source, destination, foreign, 'pdf'])
assert result.returncode != 0 and not destination.exists() and (foreign / 'keep').read_bytes() == b'keep'
isolated = root / 'no-modules'; isolated.mkdir(); shutil.copy2(exe, isolated / exe.name)
work = root / 'missing-work'; destination = root / 'missing.pdf'
result = run([isolated/exe.name, '--convert', source, destination, work, 'pdf'])
assert result.returncode != 0 and not destination.exists() and not work.exists()
assert 'Модуль' in Path(str(work) + '.error.txt').read_text('utf-8')
print('PASS presentations: corrupt/disguised/macro refusal, source preservation, exclusive output/staging, missing module', flush=True)
print(f'Fixtures: {root}', flush=True)
